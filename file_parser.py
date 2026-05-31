import os
import io
import logging
from pathlib import Path
from typing import List, Tuple, Optional
import chardet
import PyPDF2
import pdfplumber
from docx import Document
try:
    import ebooklib
    from ebooklib import epub
except ImportError:
    epub = None
    ebooklib = None
try:
    from mobi import MobiBook
except ImportError:
    MobiBook = None
from langdetect import detect, LangDetectException

logger = logging.getLogger(__name__)

class FileParser:
    """Парсер для всех поддерживаемых форматов файлов"""
    
    SUPPORTED_FORMATS = ['txt', 'pdf', 'epub', 'mobi', 'doc', 'docx']
    
    @staticmethod
    def detect_language(text: str, default: str = 'auto') -> str:
        """
        Определяет язык текста
        :param text: текст для анализа
        :param default: язык по умолчанию
        :return: код языка ('ru', 'en', 'auto')
        """
        try:
            # Берем первые 500 символов для анализа
            sample = text[:500].replace('\n', ' ')
            detected = detect(sample)
            
            if detected in ['ru', 'en']:
                return detected
            else:
                return default
        except LangDetectException:
            logger.warning("Could not detect language, using default")
            return default
    
    @staticmethod
    def parse_txt(file_path: str) -> str:
        """Парсит текстовый файл"""
        try:
            # Определяем кодировку
            with open(file_path, 'rb') as f:
                raw_data = f.read()
            
            encoding = chardet.detect(raw_data)['encoding']
            if not encoding:
                encoding = 'utf-8'
            
            with open(file_path, 'r', encoding=encoding, errors='ignore') as f:
                text = f.read()
            
            logger.info(f"Successfully parsed TXT file: {file_path} (encoding: {encoding})")
            return text
        except Exception as e:
            logger.error(f"Error parsing TXT file: {str(e)}")
            raise
    
    @staticmethod
    def parse_pdf(file_path: str) -> str:
        """Парсит PDF файл"""
        try:
            text = ""
            
            # Пытаемся использовать pdfplumber для лучшего качества
            try:
                with pdfplumber.open(file_path) as pdf:
                    for page in pdf.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text += page_text + "\n"
            except:
                # Fallback на PyPDF2
                with open(file_path, 'rb') as f:
                    reader = PyPDF2.PdfReader(f)
                    for page in reader.pages:
                        text += page.extract_text() + "\n"
            
            logger.info(f"Successfully parsed PDF file: {file_path}")
            return text
        except Exception as e:
            logger.error(f"Error parsing PDF file: {str(e)}")
            raise
    
    @staticmethod
    def parse_docx(file_path: str) -> str:
        """Парсит DOCX файл"""
        try:
            doc = Document(file_path)
            text = ""
            
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            # Добавляем текст из таблиц
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                    text += "\n"
            
            logger.info(f"Successfully parsed DOCX file: {file_path}")
            return text
        except Exception as e:
            logger.error(f"Error parsing DOCX file: {str(e)}")
            raise
    
    @staticmethod
    def parse_epub(file_path: str) -> str:
        """Парсит EPUB файл"""
        try:
            if not epub:
                raise ImportError("epub library not installed")
            
            text = ""
            book = epub.read_epub(file_path)
            
            for item in book.get_items():
                if item.get_type() == ebooklib.ITEM_DOCUMENT:
                    content = item.get_content()
                    # Простое удаление HTML тегов
                    import re
                    text_only = re.sub('<[^<]+?>', '', content.decode('utf-8', errors='ignore'))
                    text += text_only + "\n"
            
            logger.info(f"Successfully parsed EPUB file: {file_path}")
            return text
        except Exception as e:
            logger.error(f"Error parsing EPUB file: {str(e)}")
            raise
    
    @staticmethod
    def parse_mobi(file_path: str) -> str:
        """Парсит MOBI файл"""
        try:
            if not MobiBook:
                raise ImportError("mobi library not installed")
            
            tempdir = os.path.join(os.path.dirname(file_path), 'mobi_temp')
            os.makedirs(tempdir, exist_ok=True)
            
            mobi_book = MobiBook(file_path)
            mobi_book.parse()
            
            # Извлекаем текст
            text = mobi_book.content()
            
            logger.info(f"Successfully parsed MOBI file: {file_path}")
            return text
        except Exception as e:
            logger.error(f"Error parsing MOBI file: {str(e)}")
            raise
    
    @staticmethod
    def parse_doc(file_path: str) -> str:
        """Парсит DOC файл (используем python-docx если возможно)"""
        try:
            # Пытаемся как DOCX
            return FileParser.parse_docx(file_path)
        except Exception as e:
            logger.warning(f"Could not parse as DOCX, attempting alternative: {str(e)}")
            # Для старых DOC файлов нужна специальная библиотека
            raise NotImplementedError("Parsing of legacy .doc format requires python-doc2docx")
    
    @classmethod
    def parse(cls, file_path: str) -> Tuple[str, str, str]:
        """
        Универсальный парсер файла
        :param file_path: путь к файлу
        :return: (text, file_format, detected_language)
        """
        file_path = str(file_path)
        file_extension = os.path.splitext(file_path)[1].lower().strip('.')
        
        if file_extension not in cls.SUPPORTED_FORMATS:
            raise ValueError(f"Unsupported file format: {file_extension}")
        
        logger.info(f"Parsing file: {file_path} (format: {file_extension})")
        
        # Парсим файл в зависимости от формата
        if file_extension == 'txt':
            text = cls.parse_txt(file_path)
        elif file_extension == 'pdf':
            text = cls.parse_pdf(file_path)
        elif file_extension == 'docx':
            text = cls.parse_docx(file_path)
        elif file_extension == 'doc':
            text = cls.parse_doc(file_path)
        elif file_extension == 'epub':
            text = cls.parse_epub(file_path)
        elif file_extension == 'mobi':
            text = cls.parse_mobi(file_path)
        else:
            raise ValueError(f"No parser for format: {file_extension}")
        
        # Очищаем текст
        text = cls.clean_text(text)
        
        # Определяем язык
        detected_language = cls.detect_language(text)
        
        logger.info(f"Detected language: {detected_language}")
        
        return text, file_extension, detected_language
    
    @staticmethod
    def clean_text(text: str) -> str:
        """Очищает текст от лишних пробелов и символов"""
        # Удаляем множественные пробелы
        import re
        
        # Удаляем множественные новые строки
        text = re.sub(r'\n\n+', '\n\n', text)
        
        # Удаляем пробелы в начале и конце
        text = text.strip()
        
        # Заменяем множественные пробелы на один
        text = re.sub(r'  +', ' ', text)
        
        return text


if __name__ == '__main__':
    # Тестирование
    logging.basicConfig(level=logging.INFO)
    
    # Пример использования
    # text, format, lang = FileParser.parse('example.pdf')
    # print(f"Format: {format}, Language: {lang}")
    # print(f"Text length: {len(text)}")
